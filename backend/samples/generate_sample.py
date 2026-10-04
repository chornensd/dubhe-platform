import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADERS = [
    "寄件人姓名", "寄件人电话", "寄件地址", "寄件纬度", "寄件经度",
    "收件人姓名", "收件人电话", "收件地址", "收件纬度", "收件经度",
    "物品类型", "物品名称", "重量(kg)", "体积(m³)", "数量",
    "加急", "预约时间", "优惠金额", "备注",
]

wb = Workbook()
ws = wb.active
ws.title = "订单导入"
ws.append(HEADERS)
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="E8F0FE")

categories = ["文件票据", "食品", "医药用品", "电子产品", "服饰", "其他", "documents", "food"]
items = ["合同文件", "生鲜餐盒", "常规药品", "手机配件", "换洗衣物", "日用品", "发票票据", "面包糕点"]
names = ["王芳", "李强", "赵敏", "陈晨", "刘晓", "周洋", "吴迪", "郑爽"]
random.seed(42)

rows = []
for index in range(15):
    rows.append([
        random.choice(names),
        f"138{random.randint(10000000, 99999999)}",
        f"杭州市测试地址 {index + 1} 号",
        round(30.0 + random.uniform(-0.025, 0.025), 5),
        round(120.0 + random.uniform(-0.025, 0.025), 5),
        random.choice(names),
        f"139{random.randint(10000000, 99999999)}",
        f"杭州市测试地址 {index + 11} 号",
        round(30.0 + random.uniform(-0.025, 0.025), 5),
        round(120.0 + random.uniform(-0.025, 0.025), 5),
        random.choice(categories),
        random.choice(items),
        round(random.uniform(0.5, 8), 2),
        round(random.uniform(0.001, 0.05), 3),
        random.randint(1, 3),
        random.choice(["是", "否"]),
        "",
        0,
        "",
    ])

rows[14][16] = "2026-10-01 10:00"
rows[14][17] = 3

rows.append(["张三", "13700000001", "杭州市西湖区某处", 30.01, 120.01,
             "李四", "13900000002", "上海市浦东新区某处", 31.5, 121.5,
             "文件票据", "文件", 1.5, 0.01, 1, "否", "", 0, "超出服务区"])
rows.append(["张三", "13700000002", "杭州市西湖区某处", 30.01, 120.01,
             "李四", "13900000003", "杭州市滨江区某处", 30.02, 120.02,
             "易燃易爆品", "化工品", 1.0, 0.01, 1, "否", "", 0, "禁运品"])
rows.append(["张三", "123", "杭州市西湖区某处", 30.01, 120.01,
             "李四", "13900000004", "杭州市滨江区某处", 30.02, 120.02,
             "食品", "餐盒", 1.0, 0.01, 1, "否", "", 0, "电话错误"])
rows.append(["张三", "13700000004", "杭州市西湖区某处", 30.01, 120.01,
             "", "13900000005", "杭州市滨江区某处", 30.02, 120.02,
             "食品", "餐盒", 1.0, 0.01, 1, "否", "", 0, "缺收件人"])
rows.append(["张三", "13700000005", "杭州市西湖区某处", 30.01, 120.01,
             "李四", "13900000006", "杭州市滨江区某处", 30.02, 120.02,
             "电子产品", "配件", 0, 0.01, 1, "否", "", 0, "重量为0"])

for row in rows:
    ws.append(row)

for index, _ in enumerate(HEADERS, start=1):
    ws.column_dimensions[chr(64 + index)].width = 16

out_path = Path(__file__).with_name("order-import-sample.xlsx")
wb.save(out_path)
print("rows:", len(rows), "->", out_path)
